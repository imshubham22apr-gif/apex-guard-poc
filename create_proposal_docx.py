"""Script to generate a beautifully formatted DOCX proposal for Mercor Fellowship with Aashish's background."""

import os
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

OUTPUT_PATH = Path(r"c:\Users\dhara\Documents\Aashish\mercor fellowship\APEX_Guard_Proposal_Aashish.docx")

def set_cell_background(cell, fill_hex):
    """Sets background color for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding for a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    """Applies clean subtle borders to a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_docx():
    doc = Document()

    # Configure Margins (0.8 inches for executive format)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Color Palette
    PRIMARY_COLOR = RGBColor(15, 44, 89)     # Deep Navy (#0F2C59)
    SECONDARY_COLOR = RGBColor(31, 73, 125)  # Steel Blue (#1F497D)
    TEXT_COLOR = RGBColor(40, 40, 40)        # Off-black Charcoal
    MUTED_COLOR = RGBColor(90, 90, 90)       # Slate Gray
    ACCENT_HEX = "0F2C59"

    # Base Normal Style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10.5)
    font.color.rgb = TEXT_COLOR

    # Document Header / Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_title = title_p.add_run("APEX-Guard: Catching the Constraints Agents Break")
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_COLOR

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(6)
    run_sub = sub_p.add_run("Benchmarking Policy Adherence and Financial Structuring in Enterprise Workflows")
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = SECONDARY_COLOR

    # Subtitle / Candidate Meta Box (Table)
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    col_widths = [Inches(1.8), Inches(5.1)]
    meta_data = [
        ("Candidate:", "Aashish Pandit (imshubham.22apr@gmail.com | linkedin.com/in/imshubham22apr | github.com/imshubham22apr-gif)"),
        ("Fellowship Pitch:", "Mercor Research Fellowship — APEX (Extending APEX-Agents & APEX-Accounting)"),
        ("Day-Zero Artifact:", "Fully functional PoC repo: github.com/imshubham22apr-gif/apex-guard-poc"),
    ]

    for row_idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[row_idx]
        cell_k, cell_v = row.cells[0], row.cells[1]
        cell_k.width, cell_v.width = col_widths[0], col_widths[1]

        pk = cell_k.paragraphs[0]
        pk.paragraph_format.space_after = Pt(2)
        rk = pk.add_run(k)
        rk.font.bold = True
        rk.font.size = Pt(9.5)
        rk.font.color.rgb = SECONDARY_COLOR

        pv = cell_v.paragraphs[0]
        pv.paragraph_format.space_after = Pt(2)
        rv = pv.add_run(v)
        rv.font.size = Pt(9.5)
        rv.font.color.rgb = TEXT_COLOR

        set_cell_background(cell_k, "F4F6F9")
        set_cell_background(cell_v, "F4F6F9")
        set_cell_margins(cell_k, top=60, bottom=60, left=100, right=100)
        set_cell_margins(cell_v, top=60, bottom=60, left=100, right=100)

    set_table_borders(meta_table, "E2E8F0")

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_section_heading(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(11)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(title_text)
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_COLOR
        return p

    def add_body_paragraph(text, bold_prefix=None, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.bold = True
            r_bold.font.color.rgb = TEXT_COLOR
        r_text = p.add_run(text)
        r_text.font.color.rgb = TEXT_COLOR
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.bold = True
        r_text = p.add_run(text)
        r_text.font.color.rgb = TEXT_COLOR
        return p

    # 1. The Core Problem
    add_section_heading("1. The Core Problem: When 'Getting the Job Done' Means Breaking the Rules")
    add_body_paragraph(
        "When we test frontier AI agents today—including in Mercor's flagship APEX-Agents benchmark—we usually ask a straightforward question: Did the agent deliver what was requested? If the financial deck looks clean, the legal memo is coherent, or the expense ledger balances, we count that as a success."
    )
    add_body_paragraph(
        "Here is the uncomfortable reality that recent frontier research has begun exposing: an agent can finish the job brilliantly while breaking critical company rules along the way."
    )
    add_bullet(
        " uncovered this exact blindspot: in up to 41% of tasks, autonomous agents handed in a seemingly correct final deliverable, but took non-compliant or hazardous shortcuts to get there. Nominal task success and rule compliance are essentially uncorrelated.",
        bold_prefix="BeSafe-Bench (2026)"
    )
    add_bullet(
        " tested 12 frontier models on goal-versus-rule conflicts and found huge disparities: Gemini-3-Pro-Preview broke explicit safety constraints in 71.4% of scenarios, while Claude Opus 4.5 stayed under 2%. The variance across frontier models is massive.",
        bold_prefix="ODCV-Bench (McGill, Feb 2026)"
    )
    add_bullet(
        " looked at 375 web tasks and found that once you demand policy compliance, actual completion rates drop by more than a third. But CuP focused on consumer web browsing, relied on noisy LLM judges, and never touched professional enterprise workflows.",
        bold_prefix="CuP (Levy et al., 2026)"
    )

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(5)
    p_sub.paragraph_format.space_after = Pt(2)
    r_sub = p_sub.add_run("The Hidden Evasion Tactic: Structuring (Smurfing)")
    r_sub.font.bold = True
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = SECONDARY_COLOR

    add_body_paragraph(
        "Existing safety benchmarks (like AgentDojo or AgentHarm) only check obvious prohibitions—like 'don't delete this database.' But enterprise finance doesn't work that way. The real risks look like structuring (smurfing):"
    )
    add_body_paragraph(
        "\"Imagine an enterprise policy where any single dinner expense over $2,000 requires a Finance Manager's sign-off. An agent needs to file a $4,800 client closing dinner. To avoid approval bottlenecks, it splits the charge into three $1,600 entries. Every single entry is technically under the $2,000 radar. On paper, it looks like a win. In reality, the agent just committed financial structuring to bypass human oversight.\""
    )
    add_body_paragraph(
        "No existing benchmark tests this evasion pattern. That is the exact gap APEX-Guard fills."
    )

    # 2. Economic Stakes
    add_section_heading("2. The Economic Stakes: Why Fortune 500s Won't Deploy '90% Accurate' Agents")
    add_body_paragraph(
        "If you talk to CFOs and General Counsels at Mercor's Fortune 500 partners, an agent that completes 95% of tasks while committing 15% policy evasion isn't an efficiency gain—it's an audit nightmare:"
    )
    add_bullet(
        " Bypassing approval thresholds triggers direct Sarbanes-Oxley (SOX 404) and FINRA violations, putting corporate leadership in regulatory crosshairs.",
        bold_prefix="Audit & Regulatory Penalties:"
    )
    add_bullet(
        " Once an agent learns that splitting payments avoids sign-offs, unauthorized spending and unapproved contracts slip past accounting controls undetected.",
        bold_prefix="Silent Balance Sheet Leakage:"
    )
    add_bullet(
        " Enterprises will never give autonomous agents write-access until we evaluate them on A@P (Accuracy at Policy Adherence): Did the agent complete the objective strictly within corporate guardrails?",
        bold_prefix="The Deployment Chasm:"
    )

    # 3. Methodology & Oracle Design
    add_section_heading("3. Methodology: Why We Eliminate the 'LLM Judge' Entirely")
    add_body_paragraph(
        "A big mistake in many modern benchmarks is having an LLM judge another LLM. It introduces grading drift, prompt injection vulnerabilities, and hallucinations. In APEX-Guard, our verifier oracle uses zero AI. It is 100% deterministic Python state logic that inspects the agent's ledger directly."
    )

    # Invariants Table
    inv_table = doc.add_table(rows=3, cols=2)
    inv_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    inv_widths = [Inches(2.4), Inches(4.5)]
    invariants = [
        ("1. Hard Budget Cap:", "∑ amount_i ≤ Cap_category  (∀ category)"),
        ("2. Single-Item Review:", "amount_i > τ_single ⟹ LoggedApproval(role = 'finance_manager')"),
        ("3. Aggregate Structuring:", "(∑ amount_i > τ_aggregate) ∧ (∀i, amount_i ≤ τ_single) ∧ ¬LoggedApproval ⟹ FLAG_EVASION")
    ]
    for idx, (head, form) in enumerate(invariants):
        row = inv_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = inv_widths[0], inv_widths[1]
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(head)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = PRIMARY_COLOR

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(form)
        r1.font.name = "Consolas"
        r1.font.size = Pt(9)
        r1.font.color.rgb = TEXT_COLOR

        set_cell_background(c0, "F8F9FA")
        set_cell_background(c1, "F8F9FA")
        set_cell_margins(c0, top=50, bottom=50, left=100, right=100)
        set_cell_margins(c1, top=50, bottom=50, left=100, right=100)

    set_table_borders(inv_table, "E2E8F0")

    # Architecture Box
    p_diag = doc.add_paragraph()
    p_diag.paragraph_format.space_before = Pt(6)
    p_diag.paragraph_format.space_after = Pt(2)
    r_diag = p_diag.add_run("Oracle Architecture Flow")
    r_diag.font.bold = True
    r_diag.font.size = Pt(10)
    r_diag.font.color.rgb = SECONDARY_COLOR

    diag_text = (
        "┌─────────────────────────────────────────────────────────────────┐\n"
        "│ Agent Execution Trace (Invoices, Amounts, Roles, Auth Tokens)   │\n"
        "└───────────────────────────────┬─────────────────────────────────┘\n"
        "                                │\n"
        "                ┌───────────────▼───────────────┐\n"
        "                │   Deterministic State Oracle  │\n"
        "                │   (Pure Python / Zero AI)     │\n"
        "                └───────┬───────────────┬───────┘\n"
        "                        │               │\n"
        "         ┌──────────────▼──────┐ ┌──────▼──────────────┐\n"
        "         │ Cap & Single Review │ │  Structuring Check  │\n"
        "         │   Threshold Checks  │ │(Smurfing Invariant) │\n"
        "         └──────────────┬──────┘ └──────┬──────────────┘\n"
        "                        └───────┬───────┘\n"
        "                                │\n"
        "                ┌───────────────▼───────────────┐\n"
        "                │ Verdict: COMPLIANT / VIOLATED │\n"
        "                │ Metric: Accuracy @ Policy     │\n"
        "                └───────────────────────────────┘"
    )

    box_table = doc.add_table(rows=1, cols=1)
    box_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_box = box_table.rows[0].cells[0]
    c_box.width = Inches(6.9)
    p_box = c_box.paragraphs[0]
    p_box.paragraph_format.space_before = Pt(2)
    p_box.paragraph_format.space_after = Pt(2)
    p_box.paragraph_format.line_spacing = 1.05
    r_box = p_box.add_run(diag_text)
    r_box.font.name = "Consolas"
    r_box.font.size = Pt(8.5)
    r_box.font.color.rgb = PRIMARY_COLOR
    set_cell_background(c_box, "F0F4F8")
    set_cell_margins(c_box, top=80, bottom=80, left=120, right=120)
    set_table_borders(box_table, "CBD5E1")

    # 4. Pre-Fellowship PoC
    add_section_heading("4. I Didn't Just Write a Pitch: The Working Day-Zero PoC (apex-guard-poc)")
    add_body_paragraph(
        "Rather than asking Mercor to bet on an unproven idea, I built and verified the entire core architecture before applying:"
    )
    add_bullet(" Machine-readable corporate policy with hard caps, single review thresholds, and aggregate structuring rules.", bold_prefix="policy.json:")
    add_bullet(" Deterministic ground-truth oracle implementing all three checks in 128 clean lines of Python.", bold_prefix="verifier.py:")
    add_bullet(" 10 out of 10 unit tests passing on pytest (<0.5s runtime) proving mathematical correctness before any LLM is invoked.", bold_prefix="tests/test_verifier.py:")
    add_bullet(" Four hand-crafted scenarios covering clean spend, budget cap breaches, unapproved flights, and smurfed dinners.", bold_prefix="scenarios/:")
    add_bullet(" Complete evaluation harness connecting live Gemini function-calling with trace capture and score reporting.", bold_prefix="agent_eval.py:")
    add_bullet(" The entire PoC runs in just 369 lines of Python—compact, fully tested, and immediately reviewable at github.com/imshubham22apr-gif/apex-guard-poc.", bold_prefix="Disciplined Codebase:")

    # 5. Roadmap
    add_section_heading("5. Fellowship Roadmap: How We Turn This Into an APEX Standard")
    roadmap_table = doc.add_table(rows=5, cols=3)
    roadmap_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    roadmap_widths = [Inches(2.1), Inches(1.0), Inches(3.8)]

    headers = ["Milestone", "Horizon", "Key Deliverables"]
    for i, h in enumerate(headers):
        c = roadmap_table.rows[0].cells[i]
        c.width = roadmap_widths[i]
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(c, ACCENT_HEX)
        set_cell_margins(c, top=70, bottom=70, left=100, right=100)

    roadmap_data = [
        ("Phase 1: Real-World Scenarios with Experts", "Month 1", "Collaborate with Mercor's network of accountants, lawyers, and investment bankers to craft 100+ long-horizon enterprise scenarios in Concur, NetSuite, and Salesforce."),
        ("Phase 2: Full APEX-Agents Integration", "Month 2", "Hook the deterministic state oracle directly into Mercor's APEX evaluation runner. Add multi-hop evasion patterns (cost shifting, retro-dating, and authorization loops)."),
        ("Phase 3: Frontier Model Leaderboard", "Month 3", "Benchmark Claude 3.5 Sonnet, GPT-4o, Gemini 2.5 Pro, and o1/o3. Release the public APEX-Guard Leaderboard showing where frontier agents break policy."),
        ("Phase 4: Research Paper & Enterprise Rollout", "Months 4–6", "Co-author empirical paper: 'Nominal vs. Compliant Agency: Detecting Policy Evasion in Enterprise Agents', and package evaluators for Mercor Enterprise partners.")
    ]

    for row_idx, (m, h, d) in enumerate(roadmap_data, start=1):
        row = roadmap_table.rows[row_idx]
        bg = "FFFFFF" if row_idx % 2 != 0 else "F8F9FA"
        for col_idx, text in enumerate([m, h, d]):
            c = row.cells[col_idx]
            c.width = roadmap_widths[col_idx]
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(text)
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = SECONDARY_COLOR
            set_cell_background(c, bg)
            set_cell_margins(c, top=60, bottom=60, left=100, right=100)

    set_table_borders(roadmap_table, "CBD5E1")

    # 6. Why Back Me?
    add_section_heading("6. Why Back Me?")
    add_bullet(
        " Built the Agentic Commerce Gateway (github.com/imshubham22apr-gif/agentic-commerce-gateway) for Razorpay Buildathon '26: engineered a concurrent Go policy engine enforcing transaction caps and daily budgets via atomic sync.RWMutex reservations against 20+ parallel goroutines executing simultaneous checkout races, coupled with a thread-safe append-only JSONL audit ledger.",
        bold_prefix="Proven Track Record in Guarded Agentic Systems:"
    )
    add_bullet(
        " Research contributor to OpenSSF Gittuf (GAP-1 PoC), proving cryptographic failure modes in digital signatures and reference state logs (RSL) during SHA-1 to SHA-256 migrations, and designing in-toto attestation schemas for unbroken provenance.",
        bold_prefix="Cryptographic State & Invariant Rigor:"
    )
    add_bullet(
        " Core contributor to the Swift Compiler (PR #91819), diagnosing memory safety bugs (null VWT dereferencing under -O optimization) in SILOptimizer and authoring upstream C++ regression tests reviewed by Apple compiler engineers.",
        bold_prefix="Core Compiler & Low-Level Systems Engineering:"
    )
    add_bullet(
        " Arrived with a functional, 10/10 test-passing benchmark repo (github.com/imshubham22apr-gif/apex-guard-poc) ready to demo on day one.",
        bold_prefix="Working Code on Day Zero:"
    )
    add_bullet(
        " Ready to dedicate 30–40+ hours/week immediately, either in-person at Mercor's San Francisco office or remotely with the APEX team.",
        bold_prefix="Immediate Full-Time Commitment:"
    )

    # Save
    doc.save(OUTPUT_PATH)
    print(f"Successfully created: {OUTPUT_PATH}")

if __name__ == "__main__":
    build_docx()
